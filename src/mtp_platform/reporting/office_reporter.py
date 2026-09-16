"""Office 报告（xlsx / docx）—— **原生生成，不走 officecli / MCP**。

内容口径沿用基线 office_reporter（同一份 payload 派生，结论必须与 JSON 一致），
但生成方式换成 openpyxl / python-docx：

- 不再有 `create` / batch / `validate` / `save` 那套命令往返；
- xlsx **显式写列宽**（openpyxl 不自动撑列，不写就会出现 `###`）；
- docx 直接用内置 Heading 样式，不需要先补样式。

Office 报告是**附加产物**：生成失败只记 `errors`，不改变这次 run 的结论。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

_STATUS_CN = {
    "passed": "通过",
    "failed": "失败",
    "error": "错误",
    "cancelled": "取消",
    "skipped": "跳过",
    "pending": "未执行",
    "running": "执行中",
}

# 测试报告的配色：通过=绿、失败/错误=红（与「股票涨红跌绿」无关，这是报告惯例）
_FILLS = {"passed": "C6EFCE", "failed": "FFC7CE", "error": "FFC7CE", "cancelled": "FFEB9C"}
_MAX_TEXT = 300


def _require(module: str):
    try:
        return __import__(module)
    except ImportError as exc:  # pragma: no cover - 依赖缺失路径
        from mtp_contracts.errors import ConfigError

        raise ConfigError(
            f"未安装 {module}，无法生成 Office 报告",
            detail="安装： uv sync --extra office",
        ) from exc


def _clean(value: Any, *, limit: int = _MAX_TEXT) -> str:
    text = "" if value is None else str(value)
    text = " ".join(text.split())
    return text if len(text) <= limit else text[:limit] + "…"


def _counts(case: dict[str, Any]) -> str:
    counts = case.get("counts") or {}
    return f"{counts.get('assertions_passed', 0)}/{counts.get('assertions_total', 0)}"


def first_failure(case: dict[str, Any]) -> str:
    """给出该用例最早的失败点（步骤 > 断言 > 用例级错误）。"""
    for step in case.get("steps") or []:
        if step.get("status") in {"failed", "error"}:
            return f"[步骤] {step.get('step_id')}: {_clean(step.get('summary'))}"
    for assertion in case.get("assertions") or []:
        if not assertion.get("passed"):
            return f"[断言] {assertion.get('id')}: {_clean(assertion.get('message'))}"
    error = case.get("error")
    return f"[用例] {_clean(error.get('message'))}" if error else ""


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------
def _write_xlsx(payload: dict[str, Any], target: Path) -> Path:
    openpyxl = _require("openpyxl")
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    def style_header(ws, widths: list[int], row: int = 1) -> None:
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for cell in ws[row]:
            cell.font = Font(bold=True)

    # -- 汇总 ---------------------------------------------------------------
    ws = wb.active
    ws.title = "汇总"
    summary = payload["summary"]
    verdict = "PASS" if summary["success"] else "FAIL"
    rows = [
        ("run_id", payload["run_id"]),
        ("生成时间", payload.get("generated_at", "")),
        ("开始/结束", f"{payload.get('started_at', '')} ~ {payload.get('finished_at', '')}"),
        ("结论", verdict),
        ("用例 总数/通过/失败/错误/取消/跳过", (
            f"{summary['cases_total']}/{summary['cases_passed']}/{summary['cases_failed']}"
            f"/{summary['cases_error']}/{summary['cases_cancelled']}/{summary['cases_skipped']}"
        )),
        ("断言 总数/失败", f"{summary['assertions_total']}/{summary['assertions_failed']}"),
        ("总耗时(ms)", summary["duration_ms"]),
    ]
    ws.append(["项", "值"])
    for key, value in rows:
        ws.append([key, value])
    style_header(ws, [34, 60])

    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["用例 ID", "标题", "模块", "优先级", "状态", "耗时(ms)", "断言(过/总)", "首个失败"])
    for idx, width in enumerate([22, 40, 14, 10, 10, 12, 14, 60], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for case in payload["cases"]:
        row = [
            case.get("case_id"),
            _clean(case.get("title")),
            case.get("module"),
            case.get("priority"),
            _STATUS_CN.get(case.get("status"), case.get("status")),
            case.get("duration_ms"),
            _counts(case),
            first_failure(case),
        ]
        ws.append(row)
        fill = _FILLS.get(case.get("status"))
        if fill:
            ws.cell(row=ws.max_row, column=5).fill = PatternFill("solid", fgColor=fill)

    # -- 断言明细 -----------------------------------------------------------
    ws2 = wb.create_sheet("断言明细")
    ws2.append(["用例 ID", "断言 ID", "类型", "结果", "实际", "期望", "说明"])
    style_header(ws2, [22, 26, 18, 8, 40, 30, 50])
    for case in payload["cases"]:
        for assertion in case.get("assertions") or []:
            ws2.append([
                case.get("case_id"),
                assertion.get("id"),
                assertion.get("type"),
                "通过" if assertion.get("passed") else "失败",
                _clean(assertion.get("actual")),
                _clean(assertion.get("expected")),
                _clean(assertion.get("message")),
            ])
            if not assertion.get("passed"):
                ws2.cell(row=ws2.max_row, column=4).fill = PatternFill("solid", fgColor=_FILLS["failed"])

    # -- 步骤明细 -----------------------------------------------------------
    ws3 = wb.create_sheet("步骤明细")
    ws3.append(["用例 ID", "阶段", "步骤 ID", "action", "状态", "耗时(ms)", "尝试", "摘要"])
    style_header(ws3, [22, 14, 24, 26, 10, 12, 8, 60])
    for case in payload["cases"]:
        for step in case.get("steps") or []:
            ws3.append([
                case.get("case_id"),
                step.get("phase"),
                step.get("step_id"),
                step.get("action"),
                _STATUS_CN.get(step.get("status"), step.get("status")),
                step.get("duration_ms"),
                step.get("attempts"),
                _clean(step.get("summary")),
            ])
            fill = _FILLS.get(step.get("status"))
            if fill:
                ws3.cell(row=ws3.max_row, column=5).fill = PatternFill("solid", fgColor=fill)

    for sheet in (ws, ws2, ws3):
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------
def _write_docx(payload: dict[str, Any], target: Path) -> Path:
    docx = _require("docx")
    document = docx.Document()

    summary = payload["summary"]
    document.add_heading("my-test 测试报告", level=0)
    document.add_paragraph(f"run_id：{payload['run_id']}")
    document.add_paragraph(f"生成时间：{payload.get('generated_at', '')}")
    document.add_paragraph(
        f"结论：{'PASS' if summary['success'] else 'FAIL'}"
        f"（用例 {summary['cases_passed']}/{summary['cases_total']} 通过，"
        f"失败 {summary['cases_failed']}，错误 {summary['cases_error']}，"
        f"断言失败 {summary['assertions_failed']}/{summary['assertions_total']}，"
        f"耗时 {summary['duration_ms']}ms）"
    )

    document.add_heading("用例清单", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ["用例 ID", "标题", "状态", "耗时(ms)", "断言(过/总)", "首个失败"]):
        cell.text = text
    for case in payload["cases"]:
        cells = table.add_row().cells
        values = [
            str(case.get("case_id")),
            _clean(case.get("title"), limit=80),
            _STATUS_CN.get(case.get("status"), str(case.get("status"))),
            str(case.get("duration_ms")),
            _counts(case),
            _clean(first_failure(case), limit=120),
        ]
        for cell, value in zip(cells, values):
            cell.text = value

    for case in payload["cases"]:
        document.add_heading(f"{case.get('case_id')} {_clean(case.get('title'), limit=80)}", level=2)
        document.add_paragraph(
            f"状态：{_STATUS_CN.get(case.get('status'), case.get('status'))}"
            f"　耗时：{case.get('duration_ms')}ms"
            f"　模块：{case.get('module') or '-'}　优先级：{case.get('priority') or '-'}"
        )

        if case.get("steps"):
            document.add_paragraph("步骤", style="Intense Quote")
            steps = document.add_table(rows=1, cols=4)
            steps.style = "Light Grid Accent 1"
            for cell, text in zip(steps.rows[0].cells, ["阶段", "步骤 ID", "状态", "摘要"]):
                cell.text = text
            for step in case["steps"]:
                cells = steps.add_row().cells
                values = [
                    str(step.get("phase")),
                    f"{step.get('step_id')} ({step.get('action')})",
                    _STATUS_CN.get(step.get("status"), str(step.get("status"))),
                    _clean(step.get("summary"), limit=120),
                ]
                for cell, value in zip(cells, values):
                    cell.text = value

        if case.get("assertions"):
            document.add_paragraph("断言", style="Intense Quote")
            assertions = document.add_table(rows=1, cols=4)
            assertions.style = "Light Grid Accent 1"
            for cell, text in zip(assertions.rows[0].cells, ["断言 ID", "结果", "期望", "说明"]):
                cell.text = text
            for assertion in case["assertions"]:
                cells = assertions.add_row().cells
                values = [
                    str(assertion.get("id")),
                    "通过" if assertion.get("passed") else "失败",
                    _clean(assertion.get("expected"), limit=80),
                    _clean(assertion.get("message"), limit=120),
                ]
                for cell, value in zip(cells, values):
                    cell.text = value

        if case.get("evidence"):
            document.add_paragraph("证据", style="Intense Quote")
            for entry in case["evidence"]:
                if isinstance(entry, dict):
                    document.add_paragraph(
                        f"{entry.get('kind', '')} {entry.get('name', '')}: {entry.get('path', '')}",
                        style="List Bullet",
                    )

        warnings = case.get("warnings") or []
        if warnings:
            document.add_paragraph("告警", style="Intense Quote")
            for warning in warnings:
                document.add_paragraph(str(warning), style="List Bullet")

    core = document.core_properties
    core.title = f"my-test 报告 {payload['run_id']}"
    core.author = "mtp-platform"
    core.created = datetime.now()

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------
def write_office_reports(
    payload: dict[str, Any],
    out_dir: str | Path,
    *,
    filename_prefix: str = "report",
) -> dict[str, Any]:
    """生成 xlsx + docx；失败只记录，不抛出（Office 是附加产物）。"""
    out = Path(out_dir)
    result: dict[str, Any] = {"xlsx": None, "docx": None, "errors": []}

    for kind, writer, name in (
        ("xlsx", _write_xlsx, f"{filename_prefix}.xlsx"),
        ("docx", _write_docx, f"{filename_prefix}.docx"),
    ):
        try:
            result[kind] = writer(payload, out / name)
        except Exception as exc:  # noqa: BLE001 - 附加产物，失败降级
            result["errors"].append(
                {"kind": kind, "error": f"{type(exc).__name__}: {exc}", "detail": ""}
            )
    return result


__all__ = ["write_office_reports", "first_failure"]
