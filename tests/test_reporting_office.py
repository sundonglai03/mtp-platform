"""Office 报告测试（原生 xlsx / docx，用临时目录，不用 tmp_path）。"""

from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest
from docx import Document

from mtp_platform.contracts.results import CaseResult, RunState, StepResult, StepStatus
from mtp_platform.reporting import build_payload
from mtp_platform.reporting.office_reporter import first_failure, write_office_reports


def _payload() -> dict:
    ok_step = StepResult(
        step_id="s1", action="api.get", status=StepStatus.PASSED, summary="ok", duration_ms=5
    )
    bad_step = StepResult(
        step_id="s2", action="api.post", status=StepStatus.FAILED, summary="连接被拒", duration_ms=7
    )
    case = CaseResult(
        run_id="R1", case_id="C-1", title="登录接口", module="auth", priority="P1"
    )
    case.status = RunState.FAILED
    case.duration_ms = 12
    case.steps = [ok_step, bad_step]
    case.assertions = [
        {"id": "a1", "type": "equals", "passed": True, "expected": 1, "actual": 1, "message": "ok"},
        {
            "id": "a2",
            "type": "equals",
            "passed": False,
            "expected": 2,
            "actual": 3,
            "message": "值不一致",
        },
    ]
    case.evidence = [{"kind": "text", "name": "step-output", "path": "artifacts/x.txt"}]
    case.warnings = ["示例告警"]
    return build_payload([case], run_id="R1", started_at="t0", finished_at="t1")


def test_writes_both_files_without_errors():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "reports"
        result = write_office_reports(_payload(), out, filename_prefix="R1-report")

        assert result["errors"] == []
        assert Path(result["xlsx"]).is_file()
        assert Path(result["docx"]).is_file()
        assert Path(result["xlsx"]).name == "R1-report.xlsx"


def test_xlsx_summary_and_sheets():
    with tempfile.TemporaryDirectory() as tmp:
        result = write_office_reports(_payload(), Path(tmp) / "r", filename_prefix="rep")
        wb = openpyxl.load_workbook(result["xlsx"])

        assert wb.sheetnames == ["汇总", "断言明细", "步骤明细"]
        summary = wb["汇总"]
        values = [
            str(cell.value)
            for row in summary.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert any("FAIL" == v for v in values)  # 结论
        assert any("C-1" == v for v in values)  # 用例 ID

        # 列宽必须显式写过，否则 Excel 里会显示成 ###
        assert summary.column_dimensions["A"].width

        assertions = wb["断言明细"]
        assert assertions.max_row == 3  # 表头 + 2 条断言


def _all_text(doc) -> str:
    """段落 + 表格单元格里的全部文本。"""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_docx_structure_and_core_properties():
    with tempfile.TemporaryDirectory() as tmp:
        result = write_office_reports(_payload(), Path(tmp) / "r", filename_prefix="rep")
        doc = Document(result["docx"])

        text = _all_text(doc)
        assert "my-test 测试报告" in text
        assert "C-1" in text
        assert "值不一致" in text  # 失败断言说明进了断言表

        styles = {p.style.name for p in doc.paragraphs}
        assert "Title" in styles          # 报告标题（level=0）
        assert "Heading 1" in styles      # 「用例清单」
        assert "Heading 2" in styles      # 每个用例一节
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert any(h.startswith("C-1") for h in headings)

        assert "R1" in doc.core_properties.title
        assert doc.tables  # 用例清单表


def test_office_failure_is_recorded_not_raised():
    """Office 是附加产物：写不进去也要降级成告警。"""
    with tempfile.TemporaryDirectory() as tmp:
        # 用一个「已存在的文件」当输出目录，构造写入失败
        blocker = Path(tmp) / "blocked"
        blocker.write_text("not a dir", encoding="utf-8")

        result = write_office_reports(_payload(), blocker, filename_prefix="rep")

    assert result["xlsx"] is None and result["docx"] is None
    assert {e["kind"] for e in result["errors"]} == {"xlsx", "docx"}


def test_first_failure_priority_order():
    case = {
        "steps": [{"step_id": "s1", "status": "failed", "summary": "step boom"}],
        "assertions": [{"id": "a1", "passed": False, "message": "assert boom"}],
    }
    assert first_failure(case).startswith("[步骤]")
    assert first_failure({"steps": [], "assertions": case["assertions"]}).startswith("[断言]")
    assert first_failure({"error": {"message": "case boom"}}).startswith("[用例]")
    assert first_failure({}) == ""
